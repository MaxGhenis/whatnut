#!/usr/bin/env python3
"""Integrity checks over data/ (mechanical part of data/CHECK.md).

1. data/evidence.yaml parses, ids are unique, every row carries the DESIGN.md
   schema fields (top level, source, verified).
2. Every source.doi resolves on Crossref (live) and the Crossref title and first
   author match the row. Also reports year/journal/volume/issue/pages drift.
3. Every MANIFEST.json under data/: each listed file exists and its sha256
   matches; files on disk that no manifest lists are reported.
4. Life tables: the csv e0 equals the xlsx e0 (Tables 1-3) and equals e0
   recomputed from the csv qx with the mid-year convention; the csv columns
   equal the xlsx cells.
5. Cause shares: 0 <= share <= 1, cvd >= chd + stroke, cvd + cancer <= 1,
   shares equal deaths / all_deaths.

Writes a JSON summary to stdout (or --out). Rerunnable; network only for step 2.

    python scripts/check_integrity.py --out /tmp/integrity.json
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
UA = "whatnut-paper integrity check (mailto:max@maxghenis.com)"

TOP = [
    "id",
    "kind",
    "exposure",
    "outcome",
    "measure",
    "unit",
    "estimate",
    "ci_low",
    "ci_high",
    "ci_level",
    "pi_low",
    "pi_high",
    "support",
    "population",
    "notes",
    "source",
    "locator",
    "verified",
]
SRC = [
    "key",
    "doi",
    "pmid",
    "pmcid",
    "first_author",
    "authors",
    "title",
    "journal",
    "year",
    "volume",
    "issue",
    "pages",
]
VER = [
    "doi_resolves",
    "title_matches",
    "first_author_matches",
    "value_checked_in",
    "checked_on",
]
KINDS = {
    "cohort_meta_dose_response",
    "cohort",
    "rct_mediator",
    "mediator_slope",
    "rct_hard_endpoint",
    "mendelian_randomization",
    "calibration_corpus",
    "reference_value",
    "lab_or_storage_study",
    "trial_biomarker",
}
MEASURES = {"RR", "HR", "OR", "RRR", "mean_difference", "value"}
CHECKED_IN = {"abstract", "full text", "data file", "api"}
QUOTE_RE = r"\"([^\"]+)\"|“([^”]+)”|(?<![A-Za-z])'([^']{10,})'(?![A-Za-z])"


def sha256(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def toks(s) -> list[str]:
    if not s:
        return []
    s = html.unescape(re.sub(r"<[^>]+>", " ", str(s)))
    s = unicodedata.normalize("NFKD", s.lower())
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.findall(r"[a-z0-9]+", s)


# ---------------------------------------------------------------- evidence
def check_evidence() -> dict:
    path = DATA / "evidence.yaml"
    out = {"sha256": sha256(path), "rows": [], "problems": []}
    rows = yaml.safe_load(path.read_text())
    if not isinstance(rows, list):
        out["problems"].append("top level is not a list")
        return out
    seen = {}
    for i, r in enumerate(rows):
        rid = r.get("id")
        p = []
        if rid in seen:
            p.append(f"duplicate id (also row {seen[rid]})")
        seen[rid] = i
        if not rid or not re.fullmatch(r"[a-z0-9_]+", str(rid)):
            p.append("id not snake_case")
        p += [f"missing field {k}" for k in TOP if k not in r]
        p += [f"extra field {k}" for k in r if k not in TOP]
        s = r.get("source") or {}
        p += [f"missing source.{k}" for k in SRC if k not in s]
        p += [f"extra source.{k}" for k in s if k not in SRC]
        v = r.get("verified") or {}
        p += [f"missing verified.{k}" for k in VER if k not in v]
        p += [f"extra verified.{k}" for k in v if k not in VER]
        if r.get("kind") not in KINDS:
            p.append(f"kind {r.get('kind')!r} not in schema enum")
        if r.get("measure") not in MEASURES:
            p.append(f"measure {r.get('measure')!r} not in schema enum")
        if v.get("value_checked_in") not in CHECKED_IN:
            p.append(
                f"value_checked_in {v.get('value_checked_in')!r} not in schema enum"
            )
        est, lo, hi = r.get("estimate"), r.get("ci_low"), r.get("ci_high")
        if None not in (est, lo, hi) and not (lo <= est <= hi):
            p.append(f"estimate {est} outside CI [{lo}, {hi}]")
        if (lo is None) != (hi is None):
            p.append("only one CI bound set")
        if None not in (lo, hi) and r.get("ci_level") is None:
            p.append("CI set but ci_level null")
        pl, ph = r.get("pi_low"), r.get("pi_high")
        if None not in (est, pl, ph) and not (pl <= est <= ph):
            p.append(f"estimate outside PI [{pl}, {ph}]")
        if (
            r.get("measure") in ("RR", "HR", "OR", "RRR")
            and est is not None
            and est <= 0
        ):
            p.append("ratio measure <= 0")
        # long quotes: any quoted run (double, curly, or standalone single quotes)
        # inside one string field that is longer than 15 words
        for leaf in _leaves(r):
            if not isinstance(leaf, str):
                continue
            for groups in re.findall(QUOTE_RE, leaf):
                q = next(g for g in groups if g)
                if len(q.split()) > 15:
                    p.append(f"quote longer than 15 words: {q[:60]}...")
        out["rows"].append(
            {
                "id": rid,
                "kind": r.get("kind"),
                "measure": r.get("measure"),
                "estimate": est,
                "ci": [lo, hi],
                "pi": [pl, ph],
                "doi": s.get("doi"),
                "checked_in": v.get("value_checked_in"),
                "problems": p,
            }
        )
    return out


def _leaves(o):
    if isinstance(o, dict):
        for v in o.values():
            yield from _leaves(v)
    elif isinstance(o, list):
        for v in o:
            yield from _leaves(v)
    else:
        yield o


# ---------------------------------------------------------------- crossref
def crossref(doi: str) -> tuple[int, dict | None]:
    url = "https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="/()")
    for attempt in range(3):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=30) as resp:
                return resp.status, json.load(resp)["message"]
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return 404, None
            time.sleep(2 + attempt * 3)
        except Exception:
            time.sleep(2 + attempt * 3)
    return -1, None


def check_crossref() -> dict:
    rows = yaml.safe_load((DATA / "evidence.yaml").read_text())
    by_doi: dict[str, list[dict]] = {}
    for r in rows:
        d = (r.get("source") or {}).get("doi")
        if d:
            by_doi.setdefault(d.strip(), []).append(r)
    out = {}
    for doi, rs in sorted(by_doi.items()):
        status, msg = crossref(doi)
        rec = {"status": status, "rows": [r["id"] for r in rs], "issues": []}
        if msg is None:
            rec["issues"].append(f"unresolvable on Crossref (status {status})")
            out[doi] = rec
            continue
        titles = [t for t in msg.get("title") or [] if t]
        subs = [t for t in msg.get("subtitle") or [] if t]
        cands = titles + [f"{t}: {s}" for t in titles for s in subs]
        au = msg.get("author") or []
        first = [a for a in au if a.get("sequence") == "first"] or au
        cr_first = (first[0].get("family") or first[0].get("name")) if first else None
        yr = None
        for k in ("published-print", "issued", "published-online"):
            dp = (msg.get(k) or {}).get("date-parts")
            if dp and dp[0] and dp[0][0]:
                yr = dp[0][0]
                break
        rec.update(
            {
                "cr_title": cands[0] if cands else None,
                "cr_first_author": cr_first,
                "cr_journal": (msg.get("container-title") or [None])[0],
                "cr_year": yr,
                "cr_volume": msg.get("volume"),
                "cr_issue": msg.get("issue"),
                "cr_page": msg.get("page"),
            }
        )
        for r in rs:
            s = r["source"]
            lt = toks(s.get("title"))
            best_exact = any(toks(c) == lt for c in cands)
            best_j = max(
                (
                    len(set(toks(c)) & set(lt)) / max(1, len(set(toks(c)) | set(lt)))
                    for c in cands
                ),
                default=0,
            )
            if not best_exact:
                rec["issues"].append(
                    f"{r['id']}: title differs (Jaccard {best_j:.2f}); row: "
                    f"{s.get('title')!r}"
                )
            la, ca = toks(s.get("first_author")), toks(cr_first)
            if not (
                la and ca and (la == ca or set(ca) <= set(la) or set(la) <= set(ca))
            ):
                rec["issues"].append(
                    f"{r['id']}: first author row {s.get('first_author')!r} vs "
                    f"Crossref {cr_first!r}"
                )
            if s.get("year") and yr and int(s["year"]) != int(yr):
                # online-first vs print year: note, not fail
                years = set()
                for k in ("published-print", "published-online", "issued", "published"):
                    dp = (msg.get(k) or {}).get("date-parts")
                    if dp and dp[0] and dp[0][0]:
                        years.add(dp[0][0])
                if int(s["year"]) not in years:
                    rec["issues"].append(
                        f"{r['id']}: year row {s['year']} vs Crossref {sorted(years)}"
                    )
            for fld, crv in (
                ("volume", msg.get("volume")),
                ("issue", msg.get("issue")),
                ("pages", msg.get("page")),
            ):
                lv = s.get(fld)
                if lv in (None, "") or crv in (None, ""):
                    continue
                lv, crv = str(lv).strip(), str(crv).strip()
                if fld == "pages":
                    # Crossref often deposits only the first page: compare first pages
                    # then
                    lf, cf = lv.split("-")[0], crv.split("-")[0]
                    same = (
                        lv == crv
                        or ("-" not in crv and lf == cf)
                        or (
                            lf == cf
                            and len(crv.split("-")) == 2
                            and len(lv.split("-")) == 2
                            and crv.split("-")[1].endswith(lv.split("-")[1])
                        )
                    )
                else:
                    same = lv == crv
                if not same:
                    rec["issues"].append(
                        f"{r['id']}: {fld} row {lv!r} vs Crossref {crv!r}"
                    )
            jr, jc = toks(s.get("journal")), toks(rec["cr_journal"])
            if jr and jc and not (set(jr) <= set(jc) or set(jc) <= set(jr)):
                rec["issues"].append(
                    f"{r['id']}: journal row {s.get('journal')!r} vs Crossref "
                    f"{rec['cr_journal']!r}"
                )
            v = r.get("verified") or {}
            if (
                v.get("doi_resolves") is not True
                or v.get("title_matches") is not True
                or v.get("first_author_matches") is not True
            ):
                rec["issues"].append(f"{r['id']}: verified flags not all true: {v}")
        out[doi] = rec
        time.sleep(0.3)
    no_doi = [r["id"] for r in rows if not (r.get("source") or {}).get("doi")]
    return {"dois": out, "rows_without_doi": no_doi}


# ---------------------------------------------------------------- manifests
def _entries(o, path=()):
    if isinstance(o, dict):
        if "sha256" in o and isinstance(o["sha256"], str):
            yield path, o
            return
        for k, v in o.items():
            yield from _entries(v, path + (str(k),))
    elif isinstance(o, list):
        for i, v in enumerate(o):
            yield from _entries(v, path + (str(i),))


def check_manifests() -> dict:
    out = {}
    listed_all: set[Path] = set()
    for mf in sorted(DATA.rglob("MANIFEST.json")):
        base = mf.parent
        m = json.loads(mf.read_text())
        res = {
            "entries": 0,
            "ok": 0,
            "missing": [],
            "mismatch": [],
            "no_url": [],
            "no_fetched_at": [],
        }
        for path, e in _entries(m):
            rel = e.get("path") or e.get("file") or e.get("local_path") or path[-1]
            p = base / rel
            if not p.exists() and len(path) >= 2:
                p = base / "/".join(path[1:])
            res["entries"] += 1
            if not p.exists():
                res["missing"].append(rel)
                continue
            listed_all.add(p.resolve())
            got = sha256(p)
            if got != e["sha256"]:
                res["mismatch"].append(
                    {
                        "file": str(p.relative_to(base)),
                        "manifest": e["sha256"],
                        "actual": got,
                    }
                )
            else:
                res["ok"] += 1
            is_raw = (
                "raw" in p.relative_to(base).parts
                or "primary" in p.relative_to(base).parts
                or mf.parent.name in ("crossref", "texts")
            )
            if is_raw and not e.get("url") and not p.suffix == ".txt":
                res["no_url"].append(str(p.relative_to(base)))
            if is_raw and not e.get("fetched_at") and not p.suffix == ".txt":
                res["no_fetched_at"].append(str(p.relative_to(base)))
        out[str(mf.relative_to(ROOT))] = res
    unlisted = []
    for p in sorted(DATA.rglob("*")):
        if (
            p.is_file()
            and p.name not in ("MANIFEST.json", ".gitignore", "CHECK.md")
            and p.resolve() not in listed_all
            and p.name != "evidence.yaml"
            and ".DS_Store" not in p.name
        ):
            unlisted.append(str(p.relative_to(ROOT)))
    return {"manifests": out, "unlisted_files": unlisted}


# ---------------------------------------------------------------- life tables
def check_life_tables() -> dict:
    import openpyxl

    out = {}
    for sex, table in (
        ("total", "Table01"),
        ("male", "Table02"),
        ("female", "Table03"),
    ):
        p = DATA / "life_tables" / f"{sex}.csv"
        if not p.exists():
            out[sex] = {"error": "csv missing"}
            continue
        rows = list(csv.DictReader(open(p)))
        q = [float(r["qx"]) for r in rows]
        ex = [float(r["ex"]) for r in rows]
        # recompute e_x from qx: mid-year for 0..99, and open interval 100+ via the
        # table's own e100 (T100 = l100 * e100), since q100 = 1 carries no length info.
        lx = [100000.0]
        for x in range(len(q) - 1):
            lx.append(lx[-1] * (1 - q[x]))
        L = [(lx[x] + lx[x + 1]) / 2 for x in range(len(q) - 1)] + [lx[-1] * ex[-1]]
        T = [sum(L[x:]) for x in range(len(L))]
        e_mid = [T[x] / lx[x] for x in range(len(L))]
        # same, but open interval set to 0.5 years (the DESIGN's q=1 at 100, mid-year)
        L2 = L[:-1] + [lx[-1] * 0.5]
        T2 = [sum(L2[x:]) for x in range(len(L2))]
        e_mid_naive = [T2[x] / lx[x] for x in range(len(L2))]
        # xlsx cells
        wb = openpyxl.load_workbook(
            DATA / "life_tables" / "raw" / f"{table}.xlsx", data_only=True
        )
        ws = wb.active
        x_rows = []
        for row in ws.iter_rows(values_only=True):
            if (
                row
                and isinstance(row[1], (int, float))
                and isinstance(row[-1], (int, float))
                and row[0] is not None
            ):
                x_rows.append(row)
        xl_e = [float(r[-1]) for r in x_rows]
        xl_q = [float(r[1]) for r in x_rows]
        title = next(
            (c for r in ws.iter_rows(max_row=3, values_only=True) for c in r if c), None
        )
        res = {
            "n_rows": len(rows),
            "xlsx_title": title,
            "xlsx_rows": len(x_rows),
            "csv_e": {a: round(ex[a], 4) for a in (0, 20, 40, 60, 80)},
            "xlsx_e": {a: round(xl_e[a], 4) for a in (0, 20, 40, 60, 80)}
            if len(xl_e) > 80
            else None,
            "recomputed_midyear_e": {
                a: round(e_mid[a], 4) for a in (0, 20, 40, 60, 80)
            },
            "recomputed_midyear_open0.5_e": {
                a: round(e_mid_naive[a], 4) for a in (0, 20, 40, 60, 80)
            },
            "max_abs_q_diff_vs_xlsx": max(abs(a - b) for a, b in zip(q, xl_q))
            if len(xl_q) == len(q)
            else "row count differs",
            "max_abs_e_diff_vs_xlsx": max(abs(a - b) for a, b in zip(ex, xl_e))
            if len(xl_e) == len(ex)
            else "row count differs",
            "max_abs_e_diff_recomputed": max(abs(a - b) for a, b in zip(ex, e_mid)),
            "max_abs_e_diff_open0.5_ages_0_80": max(
                abs(ex[a] - e_mid_naive[a]) for a in (0, 20, 40, 60, 80)
            ),
            "ages": [int(rows[0]["age"]), int(rows[-1]["age"])],
            "q_monotone_from_10": all(q[x + 1] >= q[x] for x in range(10, len(q) - 1)),
            "last_q": q[-1],
        }
        out[sex] = res
    return out


# ---------------------------------------------------------------- causes
def check_causes() -> dict:
    out = {}
    for name in ("cause_shares.csv", "cause_shares_all_groups.csv"):
        p = DATA / "causes" / name
        if not p.exists():
            out[name] = {"error": "missing"}
            continue
        probs = []
        rows = list(csv.DictReader(open(p)))
        for r in rows:
            lab = f"{r['sex']} {r['age_group_label']}"
            n = int(r["all_deaths"])
            d = {k: int(r[f"{k}_deaths"]) for k in ("cvd", "chd", "stroke", "cancer")}
            s = {k: float(r[f"{k}_share"]) for k in ("cvd", "chd", "stroke", "cancer")}
            if d["cvd"] < d["chd"] + d["stroke"]:
                probs.append(
                    f"{lab}: cvd {d['cvd']} < chd+stroke {d['chd'] + d['stroke']}"
                )
            if d["cvd"] + d["cancer"] > n:
                probs.append(f"{lab}: cvd+cancer > all")
            for k in s:
                if not (0 <= s[k] <= 1):
                    probs.append(f"{lab}: {k}_share {s[k]} out of [0,1]")
                if abs(s[k] - d[k] / n) > 1e-6:
                    probs.append(f"{lab}: {k}_share {s[k]} != {d[k]}/{n}")
        # male + female = total
        by = {(r["sex"], r["age_group_label"]): r for r in rows}
        for (sex, lab), r in by.items():
            if sex != "total":
                continue
            m, f = by.get(("male", lab)), by.get(("female", lab))
            if m and f:
                for k in ("all", "cvd", "chd", "stroke", "cancer"):
                    if int(m[f"{k}_deaths"]) + int(f[f"{k}_deaths"]) != int(
                        r[f"{k}_deaths"]
                    ):
                        probs.append(f"{lab}: male+female {k} != total")
        tot = {
            k: sum(int(r[f"{k}_deaths"]) for r in rows if r["sex"] == "total")
            for k in ("all", "cvd", "chd", "stroke", "cancer")
        }
        out[name] = {"rows": len(rows), "problems": probs, "sum_total_rows": tot}
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out")
    ap.add_argument("--skip-crossref", action="store_true")
    a = ap.parse_args()
    res = {
        "evidence": check_evidence(),
        "manifests": check_manifests(),
        "life_tables": check_life_tables(),
        "causes": check_causes(),
    }
    if not a.skip_crossref:
        res["crossref"] = check_crossref()
    s = json.dumps(res, indent=1, default=str)
    if a.out:
        Path(a.out).write_text(s)
    else:
        print(s)
    return 0


if __name__ == "__main__":
    sys.exit(main())
