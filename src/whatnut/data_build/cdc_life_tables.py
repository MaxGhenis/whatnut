"""Fetch US 2021 life tables from CDC NVSR 72/12 Table 01.

Downloads the CDC-published Excel spreadsheet companion to NVSR Vol 72
No 12 (Arias & Xu, November 2023 — *United States Life Tables, 2021*)
once, caches it under `data/raw/cdc_nvsr/`, and writes every published
single-age qx value (ages 0-99) to `data/mortality.yaml`. Age 100 is the
CDC table's explicit terminal/open interval and is retained separately
with qx=1 rather than treated as an anchor for extrapolation.
The published remaining life expectancy for the age-100+ open interval is
also retained so survivors are not truncated to half a year.

Previously, the values in `mortality.yaml` diverged from CDC 2021 Table 1
by 20-35% in the working-age range. A reproducibility referee flagged
this as a major issue; this script eliminates the drift.

Run via: `python -m whatnut.data_build.cdc_life_tables`
(add `--apply` to overwrite mortality.yaml).
"""

from __future__ import annotations

import sys
import urllib.request
from pathlib import Path

import openpyxl

from whatnut.data_build import DATA_DIR, RAW_DIR

NVSR_URL = (
    "https://ftp.cdc.gov/pub/Health_Statistics/NCHS/Publications/"
    "NVSR/72-12/Table01.xlsx"
)
NVSR_CACHE_DIR = RAW_DIR / "cdc_nvsr"
NVSR_CACHE = NVSR_CACHE_DIR / "Table01.xlsx"

SINGLE_YEAR_AGES = tuple(range(100))
TERMINAL_AGE = 100
EXPECTED_AGES = (*SINGLE_YEAR_AGES, TERMINAL_AGE)


def download_table01() -> Path:
    """Download the CDC NVSR 72-12 Table01.xlsx and cache it locally."""
    NVSR_CACHE_DIR.mkdir(parents=True, exist_ok=True)
    if not NVSR_CACHE.exists():
        print(f"Downloading CDC Table01.xlsx from {NVSR_URL}")
        with urllib.request.urlopen(NVSR_URL) as r:
            NVSR_CACHE.write_bytes(r.read())
    return NVSR_CACHE


def parse_qx_by_age(path: Path) -> dict[int, float]:
    """Parse column B (qx) of the Table 1 sheet, keyed by start age."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Table 1"]
        qx_by_age: dict[int, float] = {}
        for row in sheet.iter_rows(values_only=True):
            age_cell, qx = row[0], row[1]
            if not isinstance(age_cell, str) or not isinstance(qx, (int, float)):
                continue
            label = age_cell.strip()
            # Formats: "0–1", "40–41", "100 and older"
            if label.startswith("100"):
                qx_by_age[100] = float(qx)
                continue
            dash = "–" if "–" in label else "-"
            head = label.split(dash, 1)[0].strip()
            if head.isdigit():
                qx_by_age[int(head)] = float(qx)
        return qx_by_age
    finally:
        workbook.close()


def parse_terminal_expected_years(path: Path) -> float:
    """Parse the published e100 value for the age-100+ open interval."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook["Table 1"]
        for row in sheet.iter_rows(values_only=True):
            age_cell = row[0]
            if (
                isinstance(age_cell, str)
                and age_cell.strip().startswith("100")
                and isinstance(row[6], (int, float))
            ):
                value = float(row[6])
                if value <= 0:
                    raise ValueError("CDC terminal life expectancy must be positive.")
                return value
    finally:
        workbook.close()
    raise ValueError("CDC life table is missing age-100 life expectancy.")


def build_mortality_yaml(
    qx_all_ages: dict[int, float],
    *,
    terminal_expected_years: float,
) -> str:
    """Return mortality YAML containing the complete CDC single-age table."""
    missing = sorted(set(EXPECTED_AGES) - set(qx_all_ages))
    if missing:
        raise ValueError(f"CDC life table is missing ages: {missing}")

    for age in SINGLE_YEAR_AGES:
        qx = qx_all_ages[age]
        if not 0 < qx < 1:
            raise ValueError(f"CDC qx at age {age} must be in (0, 1), got {qx}.")
    if qx_all_ages[TERMINAL_AGE] != 1:
        raise ValueError(
            "CDC age-100 open interval must be terminal with qx=1, "
            f"got {qx_all_ages[TERMINAL_AGE]}."
        )
    if not isinstance(terminal_expected_years, (int, float)) or not (
        0 < terminal_expected_years < 20
    ):
        raise ValueError(
            "CDC age-100 open-interval life expectancy must be in (0, 20)."
        )

    # The CDC workbook displays qx to six decimal places. Preserve those
    # published values instead of the workbook's hidden floating-point cells.
    rates = {age: round(qx_all_ages[age], 6) for age in EXPECTED_AGES}
    header = (
        "# US Life Table 2021 - mortality rates by age\n"
        "# Source: Arias E, Xu JQ. United States Life Tables, 2021.\n"
        "#   National Vital Statistics Reports, Vol 72, No 12.\n"
        "#   Hyattsville, MD: National Center for Health Statistics,\n"
        "#   November 2023.\n"
        "# https://www.cdc.gov/nchs/data/nvsr/nvsr72/nvsr72-12.pdf\n"
        "# Spreadsheet: https://ftp.cdc.gov/pub/Health_Statistics/NCHS/"
        "Publications/NVSR/72-12/Table01.xlsx\n"
        "# Raw download cached at data/raw/cdc_nvsr/Table01.xlsx.\n"
        "#\n"
        "# Values below retain every published single-age qx entry from\n"
        "# Table 1 (ages 0-99), regenerated by\n"
        "#   python -m whatnut.data_build.cdc_life_tables --apply\n"
        "# Age 100 is the table's terminal/open interval (qx=1). Its published\n"
        "# remaining life expectancy is retained for explicit 100+ person-time.\n"
        "#\n"
        "# Format: age -> annual qx (0-99); terminal open interval (100)\n"
        f"\nterminal_expected_years: {terminal_expected_years:.6f}\n"
        "\nrates:\n"
    )
    # Fixed-point decimal so YAML parses numbers as floats (not strings).
    # Scientific notation like "9e-05" is ambiguous in YAML 1.1 and can
    # round-trip as a string under safe_load.
    body = "\n".join(f"  {a}: {rates[a]:.6f}" for a in sorted(rates))
    return header + body + "\n"


def main(argv: list[str] | None = None) -> int:
    argv = argv or sys.argv[1:]
    path = download_table01()
    qx = parse_qx_by_age(path)
    terminal_expected_years = parse_terminal_expected_years(path)
    print(f"Parsed {len(qx)} ages from CDC NVSR 72-12 Table 1")
    for a in [40, 50, 60, 70, 80]:
        print(f"  qx({a}) = {qx.get(a, 'missing'):.6f}")

    yaml_text = build_mortality_yaml(
        qx,
        terminal_expected_years=terminal_expected_years,
    )
    out = DATA_DIR / "mortality.yaml"
    if "--apply" in argv:
        out.write_text(yaml_text)
        print(f"\nWrote {out}")
    else:
        print("\n--- preview of mortality.yaml ---")
        print(yaml_text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
